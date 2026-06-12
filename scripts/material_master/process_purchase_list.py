from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MATERIAL_KEYWORDS = ["物料", "材料", "名称", "品名", "商品", "货品", "产品"]
SPEC_KEYWORDS = ["规格", "型号", "材质", "尺寸", "标准", "参数"]
UNIT_KEYWORDS = ["单位", "计量"]
SUPPLIER_KEYWORDS = ["供应商", "供方", "厂家"]
DATE_KEYWORDS = ["日期", "时间"]
QTY_KEYWORDS = ["数量", "采购量"]
PRICE_KEYWORDS = ["单价", "价格", "金额"]
REMARK_KEYWORDS = ["备注", "说明", "用途"]


@dataclass
class SheetProfile:
    sheet: str
    max_row: int
    max_column: int
    header_row: int | None
    headers: list[str]
    sample_rows: list[list[Any]]


def main() -> int:
    parser = argparse.ArgumentParser(description="Profile and lighten a large purchase-list workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--out-dir", type=Path, default=Path("data/material_purchase_2024"))
    parser.add_argument("--mode", choices=["profile", "extract", "standardize"], default="profile")
    parser.add_argument("--batch-size", type=int, default=1000)
    args = parser.parse_args()

    if args.mode == "profile":
        print(json.dumps(profile_workbook(args.workbook), ensure_ascii=False, indent=2))
    elif args.mode == "extract":
        result = extract_lightweight(args.workbook, args.out_dir, args.batch_size)
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        result = standardize_lightweight(args.out_dir)
        print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def profile_workbook(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    profiles = []
    for ws in wb.worksheets:
        sample_rows = []
        header_row = None
        headers: list[str] = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [clean_cell(value) for value in row]
            if row_idx <= 10:
                sample_rows.append(values[:25])
            if header_row is None and looks_like_header(values):
                header_row = row_idx
                headers = [str(value or "").strip() for value in values]
            if row_idx >= 50 and header_row is not None:
                break
        profiles.append(
            SheetProfile(
                sheet=ws.title,
                max_row=ws.max_row,
                max_column=ws.max_column,
                header_row=header_row,
                headers=headers,
                sample_rows=sample_rows,
            ).__dict__
        )
    return profiles


def extract_lightweight(path: Path, out_dir: Path, batch_size: int) -> dict[str, Any]:
    out_dir.mkdir(parents=True, exist_ok=True)
    profiles = profile_workbook(path)
    (out_dir / "profile.json").write_text(json.dumps(profiles, ensure_ascii=False, indent=2), encoding="utf-8")

    wb = load_workbook(path, read_only=True, data_only=True)
    all_rows_path = out_dir / "purchase_rows_lite.csv"
    fieldnames = [
        "source_sheet",
        "source_row",
        "raw_material_name",
        "raw_spec",
        "unit",
        "supplier",
        "purchase_date",
        "quantity",
        "unit_price",
        "amount",
        "remark",
        "raw_row_json",
    ]
    total = 0
    by_sheet: dict[str, int] = {}
    with all_rows_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for ws in wb.worksheets:
            profile = next(item for item in profiles if item["sheet"] == ws.title)
            if not profile["header_row"]:
                continue
            header_row = int(profile["header_row"])
            headers = profile["headers"]
            mapping = infer_mapping(headers)
            sheet_count = 0
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if row_idx <= header_row:
                    continue
                values = [clean_cell(value) for value in row]
                record = row_to_record(ws.title, row_idx, headers, values, mapping)
                if not record["raw_material_name"] and not record["raw_spec"]:
                    continue
                writer.writerow(record)
                total += 1
                sheet_count += 1
            by_sheet[ws.title] = sheet_count

    batch_paths = split_csv(all_rows_path, out_dir, batch_size)
    return {
        "profile": str(out_dir / "profile.json"),
        "lite_csv": str(all_rows_path),
        "batch_count": len(batch_paths),
        "batches": [str(path) for path in batch_paths],
        "row_count": total,
        "rows_by_sheet": by_sheet,
    }


def standardize_lightweight(out_dir: Path) -> dict[str, Any]:
    source = out_dir / "purchase_rows_lite.csv"
    if not source.exists():
        raise FileNotFoundError(source)

    groups: dict[tuple[str, str, str], dict[str, Any]] = {}
    aliases: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            parsed = parse_material(row["raw_material_name"], row["raw_spec"])
            key = (
                parsed["standard_name"],
                parsed["specification"],
                normalize_unit(row["unit"]),
            )
            if not key[0] and not key[1]:
                continue
            item = groups.setdefault(
                key,
                {
                    "standard_name": key[0],
                    "specification": key[1],
                    "unit": key[2],
                    "item_group_key": parsed["item_group_key"],
                    "erpnext_item_group": parsed["erpnext_item_group"],
                    "material": parsed.get("material", ""),
                    "thickness": parsed.get("thickness", ""),
                    "width": parsed.get("width", ""),
                    "form": parsed.get("form", ""),
                    "brand": parsed.get("brand", ""),
                    "model": parsed.get("model", ""),
                    "package_spec": parsed.get("package_spec", ""),
                    "source_count": 0,
                    "suppliers": Counter(),
                    "source_rows": [],
                    "status": parsed["status"],
                    "notes": parsed["notes"],
                },
            )
            item["source_count"] += 1
            if row["supplier"]:
                item["suppliers"][row["supplier"]] += 1
            item["source_rows"].append(f"{row['source_sheet']}:{row['source_row']}")
            aliases[key].add(row["raw_material_name"])

    standard_path = out_dir / "standard_material_candidates.csv"
    alias_path = out_dir / "material_alias_candidates.csv"
    unresolved_path = out_dir / "material_unresolved.csv"

    standard_fields = [
        "proposed_item_code",
        "standard_name",
        "specification",
        "unit",
        "item_group_key",
        "erpnext_item_group",
        "material",
        "thickness",
        "width",
        "form",
        "brand",
        "model",
        "package_spec",
        "alias_names",
        "supplier_names",
        "source_count",
        "status",
        "notes",
        "source_rows",
    ]
    counters: Counter[str] = Counter()
    unresolved = []
    with standard_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=standard_fields)
        writer.writeheader()
        for key, item in sorted(groups.items(), key=lambda pair: (-pair[1]["source_count"], pair[0])):
            prefix = code_prefix_for(item["item_group_key"])
            counters[prefix] += 1
            row = dict(item)
            row["proposed_item_code"] = f"{prefix}-{counters[prefix]:06d}"
            row["alias_names"] = " | ".join(sorted(aliases[key]))
            row["supplier_names"] = " | ".join(name for name, _ in item["suppliers"].most_common(8))
            row["source_rows"] = " | ".join(item["source_rows"][:20])
            row.pop("suppliers")
            writer.writerow(row)
            if item["status"] != "ready":
                unresolved.append(row)

    with alias_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["raw_name", "standard_name", "specification", "unit", "item_group_key"])
        writer.writeheader()
        for key, names in sorted(aliases.items()):
            for name in sorted(names):
                writer.writerow(
                    {
                        "raw_name": name,
                        "standard_name": key[0],
                        "specification": key[1],
                        "unit": key[2],
                        "item_group_key": groups[key]["item_group_key"],
                    }
                )

    with unresolved_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=standard_fields)
        writer.writeheader()
        writer.writerows(unresolved)

    return {
        "standard_material_candidates": str(standard_path),
        "material_alias_candidates": str(alias_path),
        "material_unresolved": str(unresolved_path),
        "standard_count": len(groups),
        "unresolved_count": len(unresolved),
    }


def looks_like_header(values: list[Any]) -> bool:
    text = " ".join(str(value or "") for value in values)
    groups = [MATERIAL_KEYWORDS, SPEC_KEYWORDS, UNIT_KEYWORDS, SUPPLIER_KEYWORDS, QTY_KEYWORDS]
    return sum(any(keyword in text for keyword in group) for group in groups) >= 2


def infer_mapping(headers: list[str]) -> dict[str, int | None]:
    return {
        "raw_material_name": find_header(headers, MATERIAL_KEYWORDS),
        "raw_spec": find_header(headers, SPEC_KEYWORDS),
        "unit": find_header(headers, UNIT_KEYWORDS),
        "supplier": find_header(headers, SUPPLIER_KEYWORDS),
        "purchase_date": find_header(headers, DATE_KEYWORDS),
        "quantity": find_header(headers, QTY_KEYWORDS),
        "unit_price": find_header(headers, ["单价"]),
        "amount": find_header(headers, ["金额", "价税", "合计"]),
        "remark": find_header(headers, REMARK_KEYWORDS),
    }


def find_header(headers: list[str], keywords: list[str]) -> int | None:
    for idx, header in enumerate(headers):
        if any(keyword in str(header) for keyword in keywords):
            return idx
    return None


def row_to_record(sheet: str, row_idx: int, headers: list[str], values: list[Any], mapping: dict[str, int | None]) -> dict[str, Any]:
    raw_map = {headers[idx] if idx < len(headers) else f"col_{idx + 1}": values[idx] for idx in range(min(len(headers), len(values)))}
    record = {
        "source_sheet": sheet,
        "source_row": row_idx,
        "raw_material_name": get_mapped(values, mapping["raw_material_name"]),
        "raw_spec": get_mapped(values, mapping["raw_spec"]),
        "unit": get_mapped(values, mapping["unit"]),
        "supplier": get_mapped(values, mapping["supplier"]),
        "purchase_date": get_mapped(values, mapping["purchase_date"]),
        "quantity": get_mapped(values, mapping["quantity"]),
        "unit_price": get_mapped(values, mapping["unit_price"]),
        "amount": get_mapped(values, mapping["amount"]),
        "remark": get_mapped(values, mapping["remark"]),
        "raw_row_json": json.dumps(raw_map, ensure_ascii=False),
    }
    return {key: "" if value is None else value for key, value in record.items()}


def split_csv(path: Path, out_dir: Path, batch_size: int) -> list[Path]:
    batches_dir = out_dir / "batches"
    batches_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        writer = None
        output = None
        for idx, row in enumerate(reader, start=1):
            batch_no = (idx - 1) // batch_size + 1
            if (idx - 1) % batch_size == 0:
                if output:
                    output.close()
                batch_path = batches_dir / f"purchase_rows_lite_batch_{batch_no:04d}.csv"
                paths.append(batch_path)
                output = batch_path.open("w", encoding="utf-8-sig", newline="")
                writer = csv.DictWriter(output, fieldnames=reader.fieldnames or [])
                writer.writeheader()
            assert writer is not None
            writer.writerow(row)
        if output:
            output.close()
    return paths


def parse_material(raw_name: str, raw_spec: str) -> dict[str, str]:
    text = f"{raw_name} {raw_spec}".strip()
    normalized = normalize_text(text)
    result = {
        "standard_name": raw_name.strip(),
        "specification": raw_spec.strip(),
        "item_group_key": "consumable",
        "erpnext_item_group": "耗材",
        "status": "needs_review",
        "notes": "规则未能完全识别，请人工确认分类和规格。",
    }

    if any(token in normalized for token in ["冷板", "冷轧", "钢板", "铁板", "卷料", "板材"]):
        result.update({"standard_name": "冷轧钢卷" if "卷" in normalized else "冷轧钢板", "item_group_key": "raw_metal_sheet", "erpnext_item_group": "原材料"})
        result["material"] = first_match(normalized, [r"SPCC", r"DC01", r"SECC", r"SGCC", r"Q235", r"Q345"])
        result["thickness"] = first_match(normalized, [r"(\d+(?:\.\d+)?)\s*(?:mm|毫米)?\s*(?:厚|厚度)?"])
        result["width"] = first_match(normalized, [r"(?:宽|宽度)?\s*(\d{3,4})\s*(?:mm|毫米)?"])
        result["form"] = "卷料" if "卷" in normalized else ("板料" if "板" in normalized else "")
        parts = [result.get("material", ""), result.get("thickness", ""), result.get("width", ""), result.get("form", "")]
        result["specification"] = " ".join(part for part in parts if part) or raw_spec.strip()
        result["status"] = "ready" if result.get("material") and result.get("thickness") and result.get("width") else "needs_review"
        result["notes"] = "" if result["status"] == "ready" else "金属板材缺少材质/厚度/宽度等关键规格。"
    elif any(token in normalized for token in ["纸箱", "包装", "胶袋", "标签"]):
        result.update({"item_group_key": "packaging_material", "erpnext_item_group": "包装材料", "status": "needs_review", "notes": "包装材料需确认材质、尺寸、包装规格。"})
    elif any(token in normalized for token in ["电阻", "电容", "芯片", "IC", "元件"]):
        result.update({"item_group_key": "electronic_component", "erpnext_item_group": "原材料", "status": "needs_review", "notes": "电子元器件需确认型号、品牌、封装、关键参数。"})
    elif any(token in normalized for token in ["服务", "安装", "维修"]):
        result.update({"item_group_key": "service_item", "erpnext_item_group": "服务", "status": "needs_review", "notes": "服务项需确认服务范围和计价单位。"})

    if not result["specification"]:
        result["specification"] = normalized
    return result


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("*", "x").replace("×", "x")).strip()


def first_match(text: str, patterns: list[str]) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return match.group(1) if match.groups() else match.group(0).upper()
    return ""


def normalize_unit(value: str) -> str:
    value = str(value or "").strip()
    return {"公斤": "Kg", "千克": "Kg", "KG": "Kg", "kg": "Kg", "个": "Nos", "件": "Nos", "箱": "Box"}.get(value, value)


def code_prefix_for(group_key: str) -> str:
    return {
        "raw_metal_sheet": "RM-MET-SHT",
        "raw_plastic_granule": "RM-PLA-GRN",
        "electronic_component": "RM-ELE-CMP",
        "packaging_material": "PKG",
        "consumable": "CNS",
        "spare_part": "SP",
        "semi_finished_good": "SF",
        "finished_good": "FG",
        "service_item": "SV",
    }.get(group_key, "MAT")


def get_mapped(values: list[Any], idx: int | None) -> Any:
    if idx is None or idx >= len(values):
        return ""
    return values[idx]


def clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, str):
        return re.sub(r"\s+", " ", value).strip()
    return value


if __name__ == "__main__":
    raise SystemExit(main())
